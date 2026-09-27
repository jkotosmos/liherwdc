"""Готовые документы файлом: КП, отчёт, протокол — в Word, расчёт — в Excel.

ТЗ требует «готовить сводки, отчёты, расчёты, коммерческие предложения и
проекты документов». Текст в чате — черновик, а работать дальше удобно с
файлом: переслать, поправить, распечатать. Инструмент собирает .docx или
.xlsx из размеченного текста и отдаёт его пользователю в чат.

Подтверждения не требует: файл получает только тот, кто попросил, во внешнем
мире ничего не меняется. Положить его на Google Диск — отдельное действие
(drive_create_file) и уже с подтверждением.

Разметка простая: «# » «## » «### » — заголовки, «- » — пункты, «1. » —
нумерация, строки «| a | b |» — таблица, **жирный**.
"""

from __future__ import annotations

import re
import secrets
import time
from pathlib import Path
from typing import Any

from ..config import settings
from ..storage import read_json, write_json
from .base import ToolError, ToolSpec, registry

KEEP_SECONDS = 7 * 24 * 3600
MAX_CONTENT = 200_000
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _dir() -> Path:
    return settings.data_dir / "outbox"


def _index_path() -> Path:
    return _dir() / "index.json"


def _safe_name(name: str, extension: str) -> str:
    stem = re.sub(r"[^\w\s.-]", "", name or "", flags=re.UNICODE).strip().strip(".")
    stem = re.sub(r"\s+", " ", stem)[:80] or "Документ"
    if stem.lower().endswith("." + extension):
        stem = stem[: -len(extension) - 1]
    return f"{stem}.{extension}"


# --- разбор разметки ----------------------------------------------------------

_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def blocks(text: str) -> list[tuple[str, Any]]:
    """Текст → [(вид, данные)]: heading1-3, bullet, number, table, paragraph."""
    result: list[tuple[str, Any]] = []
    lines = (text or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if _TABLE_ROW.match(line):
            rows = []
            while i < len(lines) and _TABLE_ROW.match(lines[i]):
                if not _TABLE_SEP.match(lines[i]):
                    rows.append(_cells(lines[i]))
                i += 1
            result.append(("table", rows))
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            result.append((f"heading{len(heading.group(1))}", heading.group(2).strip()))
        elif re.match(r"^[-*•]\s+", stripped):
            result.append(("bullet", re.sub(r"^[-*•]\s+", "", stripped)))
        elif re.match(r"^\d+[.)]\s+", stripped):
            result.append(("number", re.sub(r"^\d+[.)]\s+", "", stripped)))
        else:
            result.append(("paragraph", stripped))
        i += 1
    return result


def _runs(paragraph, text: str) -> None:
    """**жирный** внутри строки."""
    for index, part in enumerate(re.split(r"\*\*", text)):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1


def _plain(text: str) -> str:
    return text.replace("**", "")


def build_docx(title: str, content: str, path: Path) -> None:
    from docx import Document

    document = Document()
    if title:
        document.add_heading(title, level=0)
    for kind, data in blocks(content):
        if kind.startswith("heading"):
            document.add_heading(_plain(data), level=int(kind[-1]))
        elif kind == "bullet":
            _runs(document.add_paragraph(style="List Bullet"), data)
        elif kind == "number":
            _runs(document.add_paragraph(style="List Number"), data)
        elif kind == "table":
            width = max(len(row) for row in data)
            table = document.add_table(rows=len(data), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(data):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        else:
            _runs(document.add_paragraph(), data)
    document.save(path)


def _cell_value(text: str) -> Any:
    """Число — числом, чтобы в Excel работали формулы и суммы."""
    raw = _plain(text).replace(" ", " ").strip()
    candidate = raw.replace(" ", "").replace(",", ".").rstrip("%").rstrip("₽").strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", candidate):
        number = float(candidate)
        if raw.endswith("%"):
            return number / 100
        return int(number) if number.is_integer() else number
    return raw


def build_xlsx(title: str, content: str, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    workbook.remove(workbook.active)
    tables = [data for kind, data in blocks(content) if kind == "table"]
    if not tables:
        # Таблиц нет — строки текста в столбец: лучше, чем пустой файл.
        tables = [[[line] for line in (content or "").splitlines() if line.strip()]]
    for number, rows in enumerate(tables, 1):
        sheet = workbook.create_sheet(title=(title or "Лист")[:25] + (f" {number}" if len(tables) > 1 else ""))
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                cell = sheet.cell(row=r, column=c, value=_cell_value(value))
                if r == 1 and len(rows) > 1:
                    cell.font = Font(bold=True)
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 60)
    workbook.save(path)


# --- хранение и выдача -------------------------------------------------------


def _cleanup(index: dict) -> None:
    now = time.time()
    for file_id, meta in list(index.items()):
        if now - meta.get("created", 0) > KEEP_SECONDS:
            Path(meta.get("path", "")).unlink(missing_ok=True)
            index.pop(file_id, None)


def get(file_id: str) -> dict | None:
    meta = read_json(_index_path(), {}).get(file_id)
    if not meta or not Path(meta.get("path", "")).is_file():
        return None
    return meta


def _prepare(tool_input: dict[str, Any]) -> Any:
    fmt = (tool_input.get("format") or "docx").strip().lower()
    if fmt not in MIME:
        raise ToolError("format: docx (документ) или xlsx (таблица).")
    title = (tool_input.get("title") or "").strip()
    content = tool_input.get("content") or ""
    if not content.strip():
        raise ToolError("Пустое содержимое (content).")
    if len(content) > MAX_CONTENT:
        raise ToolError("Документ слишком большой — разделите на части.")

    name = _safe_name(tool_input.get("filename") or title, fmt)
    _dir().mkdir(parents=True, exist_ok=True)
    file_id = secrets.token_urlsafe(18)
    path = _dir() / f"{file_id}.{fmt}"
    (build_docx if fmt == "docx" else build_xlsx)(title, content, path)

    index = read_json(_index_path(), {})
    _cleanup(index)
    index[file_id] = {"path": str(path), "name": name, "mime": MIME[fmt], "created": time.time()}
    write_json(_index_path(), index)

    return {
        "status": "ready",
        "attachment": {"id": file_id, "name": name, "size": path.stat().st_size},
        "note": (
            f"Файл «{name}» подготовлен и будет отправлен пользователю в чат. "
            "Не пересказывай содержимое целиком — коротко скажи, что в файле."
        ),
    }


registry.register(
    ToolSpec(
        name="document_prepare",
        description=(
            "Готовит файл и отправляет его пользователю в чат: Word (.docx) — коммерческое "
            "предложение, отчёт, сводка, протокол, проект договора или письма; Excel (.xlsx) — "
            "расчёт, таблица показателей. Разметка content: «# », «## », «### » — заголовки, "
            "«- » — пункты, «1. » — нумерация, «| a | b |» — таблица (первая строка — шапка), "
            "**жирный**. Цифры берутся только из источников и инструмента calculate — не выдумывай. "
            "Файл получает только пользователь; на Google Диск — отдельно через drive_create_file."
        ),
        input_schema={
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Заголовок документа."},
                "content": {"type": "string", "description": "Текст с разметкой."},
                "format": {"type": "string", "enum": ["docx", "xlsx"], "description": "По умолчанию docx."},
                "filename": {"type": "string", "description": "Имя файла без пути (необязательно)."},
            },
            "required": ["title", "content"],
        },
        handler=_prepare,
        activity="Готовлю документ",
    )
)
