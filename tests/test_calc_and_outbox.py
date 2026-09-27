"""Точные расчёты и готовые документы файлом."""

from __future__ import annotations

import json

import pytest
from fastapi.testclient import TestClient

from app.errors import ToolError
from app.tools import calc, outbox, registry


def run(tool: str, payload: dict):
    content, is_error = registry.execute(tool, payload)
    return content, is_error


class TestCalculate:
    @pytest.mark.parametrize(("expression", "value"), [
        ("1200*15", 18000),
        ("1 250 000,50 * 0.2", 250000.1),
        ("12,5 * 2", 25),
        ("pct(25,200)", 12.5),
        ("change(110,100)", 10),
        ("avg([1,2,3])", 2),
        ("round(2/3, 2)", 0.67),
        ("max(3, 7.5, 2)", 7.5),
    ])
    def test_values(self, expression, value) -> None:
        assert calc.evaluate(expression) == pytest.approx(value)

    def test_named_steps_for_an_offer(self) -> None:
        content, is_error = run("calculate", {
            "variables": {"price": "1 200", "qty": 15},
            "steps": [
                {"name": "sum", "expression": "price*qty"},
                {"name": "discount", "expression": "sum*0.1"},
                {"name": "total", "expression": "sum-discount"},
            ],
        })
        assert not is_error
        results = {r["name"]: r for r in json.loads(content)["results"]}
        assert results["total"]["value"] == pytest.approx(16200)
        assert results["total"]["formatted"] == "16 200"

    @pytest.mark.parametrize("expression", ['__import__("os")', "open('x')", "(1).real", "2**1000", "x+1"])
    def test_unsafe_or_unknown_rejected(self, expression) -> None:
        with pytest.raises(ToolError):
            calc.evaluate(expression)

    def test_division_by_zero_named(self) -> None:
        with pytest.raises(ToolError, match="Деление на ноль"):
            calc.evaluate("5/0")

    def test_not_gated(self) -> None:
        assert registry.get("calculate").requires_confirmation is False


OFFER = """# Коммерческое предложение
Для **ООО Ромашка**.

## Состав
- Платформа OPERON
- Внедрение

| Позиция | Кол-во | Цена |
|---|---|---|
| Лицензия | 15 | 1 200 |
| Внедрение | 1 | 50 000,50 |

1. Срок действия — 30 дней
"""


class TestDocuments:
    def test_docx_structure(self) -> None:
        from docx import Document

        content, is_error = run("document_prepare", {"title": "КП Ромашка", "content": OFFER})
        assert not is_error
        meta = outbox.get(json.loads(content)["attachment"]["id"])
        assert meta["name"] == "КП Ромашка.docx"
        document = Document(meta["path"])
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Коммерческое предложение" in text and "Платформа OPERON" in text
        assert any(run.bold and "Ромашка" in run.text for p in document.paragraphs for run in p.runs)
        table = document.tables[0]
        assert [c.text for c in table.rows[1].cells] == ["Лицензия", "15", "1 200"]

    def test_xlsx_numbers_are_numbers(self) -> None:
        from openpyxl import load_workbook

        content, is_error = run("document_prepare", {"title": "Расчёт", "content": OFFER, "format": "xlsx"})
        assert not is_error
        meta = outbox.get(json.loads(content)["attachment"]["id"])
        sheet = load_workbook(meta["path"]).active
        assert sheet["A1"].value == "Позиция"
        assert sheet["C2"].value == 1200, "«1 200» должно стать числом, чтобы работали формулы"
        assert sheet["C3"].value == pytest.approx(50000.5)

    def test_filename_is_sanitized(self) -> None:
        content, _ = run("document_prepare", {"title": "x", "content": "текст", "filename": "../../etc/passwd"})
        name = json.loads(content)["attachment"]["name"]
        assert "/" not in name and name.endswith(".docx")

    def test_not_gated_it_changes_nothing_outside(self) -> None:
        assert registry.get("document_prepare").requires_confirmation is False

    def test_download_endpoint(self) -> None:
        from app import server

        content, _ = run("document_prepare", {"title": "Отчёт", "content": "# Итоги\nВсё хорошо"})
        file_id = json.loads(content)["attachment"]["id"]
        with TestClient(server.app) as client:
            response = client.get(f"/api/files/{file_id}")
            assert response.status_code == 200
            assert response.content[:2] == b"PK", "docx — это zip"
            assert client.get("/api/files/нет-такого").status_code == 404


def test_agent_emits_attachment_event() -> None:
    from fakes import FakeClient, text_block, tool_use_block, turn

    from app.agent import OperonAgent, Session

    agent = OperonAgent()
    agent._client = FakeClient([
        turn(tool_use_block("t1", "document_prepare", {"title": "КП", "content": "# КП\nтекст"}), stop_reason="tool_use"),
        turn(text_block("Готово, КП в файле.")),
    ])
    events = list(agent.send_user_message(Session(session_id="doc"), "подготовь КП"))
    attachments = [e for e in events if e["type"] == "attachment"]
    assert len(attachments) == 1 and attachments[0]["name"] == "КП.docx"


def test_telegram_send_document_is_multipart(tmp_path) -> None:
    import httpx

    from app.telegram.api import TelegramAPI

    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["path"] = request.url.path
        seen["type"] = request.headers["content-type"]
        seen["body"] = request.read()
        return httpx.Response(200, json={"ok": True, "result": {"message_id": 5}})

    file = tmp_path / "КП.docx"
    file.write_bytes(b"PK\x03\x04docx")
    api = TelegramAPI("123:abc")
    api._client = httpx.Client(transport=httpx.MockTransport(handler))
    api.send_document(42, str(file), "КП.docx")
    assert seen["path"].endswith("/sendDocument")
    assert seen["type"].startswith("multipart/form-data")
    assert b"PK\x03\x04docx" in seen["body"] and b'name="chat_id"' in seen["body"]
