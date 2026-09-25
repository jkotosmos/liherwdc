"""Офисные форматы: договоры лежат в .docx и .pdf, а не в markdown.

Файлы здесь не подделываются словарями — они собираются настоящими
библиотеками и разбираются настоящим кодом. Иначе тест проверял бы только
собственные допущения о формате.
"""

from __future__ import annotations

import io

import pytest

from app.kb.documents import (
    DOCUMENT_SUFFIXES,
    DocumentError,
    extract_text,
    is_supported,
    legacy_alternative,
)


# --- настоящие файлы -------------------------------------------------------


def make_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Договор поставки № 17/2026")
    document.add_paragraph("Срок оплаты — 30 календарных дней.")
    table = document.add_table(rows=2, cols=3)
    for col, value in enumerate(("Позиция", "Цена", "Срок")):
        table.cell(0, col).text = value
    for col, value in enumerate(("Монтаж", "450 000 ₽", "45 дней")):
        table.cell(1, col).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Тарифы"
    sheet.append(["Тариф", "Абонплата", "Скидка"])
    sheet.append(["Базовый", 12000, "0%"])
    sheet.append(["Расширенный", 28000, "15%"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Итоги квартала"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def make_pdf(text: str = "Регламент согласования скидок") -> bytes:
    from pypdf import PdfWriter

    # Пустая страница: текстового слоя нет — ровно случай скана.
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class TestWord:
    def test_paragraphs_are_extracted(self) -> None:
        text = extract_text(make_docx(), "договор.docx")
        assert "Договор поставки № 17/2026" in text
        assert "30 календарных дней" in text

    def test_tables_become_searchable_rows(self) -> None:
        """Ячейка на строку бесполезна для поиска: смысл в строке целиком."""
        text = extract_text(make_docx(), "договор.docx")
        assert "Позиция | Цена | Срок" in text
        assert "Монтаж | 450 000 ₽ | 45 дней" in text

    def test_broken_file_says_why(self) -> None:
        with pytest.raises(DocumentError, match="не читается"):
            extract_text("это не docx".encode(), "битый.docx")


class TestExcel:
    def test_sheets_and_rows(self) -> None:
        text = extract_text(make_xlsx(), "тарифы.xlsx")
        assert "Лист: Тарифы" in text
        assert "Базовый | 12000 | 0%" in text
        assert "Расширенный | 28000 | 15%" in text

    def test_empty_book_is_reported(self) -> None:
        from openpyxl import Workbook

        buffer = io.BytesIO()
        Workbook().save(buffer)
        with pytest.raises(DocumentError, match="нет заполненных"):
            extract_text(buffer.getvalue(), "пусто.xlsx")


class TestPowerPoint:
    def test_slide_text(self) -> None:
        text = extract_text(make_pptx(), "итоги.pptx")
        assert "Слайд 1" in text
        assert "Итоги квартала" in text


class TestPdf:
    def test_scan_without_text_layer_is_explained(self) -> None:
        """Молча вернуть пустоту — значит получить уверенный неверный ответ."""
        with pytest.raises(DocumentError, match="скан"):
            extract_text(make_pdf(), "скан.pdf")

    def test_garbage_is_not_silently_empty(self) -> None:
        with pytest.raises(DocumentError):
            extract_text("%PDF-1.4 мусор".encode(), "битый.pdf")


class TestFormatRouting:
    @pytest.mark.parametrize("name", ["а.docx", "б.pdf", "в.xlsx", "г.pptx", "д.XLSM"])
    def test_supported(self, name: str) -> None:
        assert is_supported(name) is True

    @pytest.mark.parametrize("name", ["а.md", "б.jpg", "в.zip", "г.doc"])
    def test_not_supported(self, name: str) -> None:
        assert is_supported(name) is False

    def test_legacy_office_suggests_the_new_format(self) -> None:
        """«Файл пуст» — худший ответ: пользователь решит, что документ плохой."""
        assert legacy_alternative("старый.doc") == ".docx"
        assert legacy_alternative("старый.xls") == ".xlsx"
        with pytest.raises(DocumentError, match="Пересохраните файл как .docx"):
            extract_text(b"", "старый.doc")

    def test_unknown_format_is_named(self) -> None:
        with pytest.raises(DocumentError, match=r"\.zip"):
            extract_text(b"", "архив.zip")


class TestKnowledgeBaseIndexing:
    """Сквозной путь: файл на диске → индекс → поиск."""

    def test_docx_becomes_searchable(self, tmp_path, monkeypatch) -> None:
        from app.kb.store import KnowledgeBase

        (tmp_path / "договоры").mkdir()
        (tmp_path / "договоры" / "поставка.docx").write_bytes(make_docx())

        kb = KnowledgeBase(root=tmp_path)
        kb.ensure_fresh()

        assert kb.stats["documents"] == 1
        results = kb.search("срок оплаты")
        assert results, "документ .docx должен находиться поиском"
        assert "поставка" in results[0]["doc_id"] or "поставка" in results[0]["title"]

    def test_xlsx_numbers_are_searchable(self, tmp_path) -> None:
        from app.kb.store import KnowledgeBase

        (tmp_path / "тарифы.xlsx").write_bytes(make_xlsx())
        kb = KnowledgeBase(root=tmp_path)
        kb.ensure_fresh()

        results = kb.search("расширенный тариф")
        assert results
        assert any("28000" in r["text"] for r in results)

    def test_broken_document_does_not_break_the_index(self, tmp_path) -> None:
        """Один битый файл не должен лишать доступа ко всей базе."""
        from app.kb.store import KnowledgeBase

        (tmp_path / "битый.docx").write_bytes("не docx".encode())
        (tmp_path / "хороший.md").write_text("# Регламент\n\nСкидки согласует директор.", encoding="utf-8")

        kb = KnowledgeBase(root=tmp_path)
        kb.ensure_fresh()

        assert kb.search("скидки"), "исправный документ обязан остаться доступным"

    def test_office_suffixes_are_indexed(self) -> None:
        from app.kb.store import SUPPORTED_SUFFIXES

        assert DOCUMENT_SUFFIXES <= SUPPORTED_SUFFIXES


class TestDriveReading:
    """drive_read раньше отвечал отказом на PDF и Word."""

    def test_office_mime_types_are_mapped(self) -> None:
        from app.tools.drive import OFFICE_MIME_SUFFIX

        assert OFFICE_MIME_SUFFIX["application/pdf"] == ".pdf"
        assert (
            OFFICE_MIME_SUFFIX[
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ]
            == ".docx"
        )

    def test_downloaded_docx_is_parsed(self, monkeypatch) -> None:
        from app.tools import drive as drive_tools

        class FakeFiles:
            def get_media(self, fileId, supportsAllDrives=False):
                return self

            def execute(self):
                return make_docx()

        class FakeDrive:
            def files(self):
                return FakeFiles()

        monkeypatch.setattr(drive_tools, "_drive", lambda: FakeDrive())
        text = drive_tools._export_text(
            "id",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "договор.docx",
        )
        assert "Договор поставки № 17/2026" in text
